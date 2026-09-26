"""Read-only source extraction. Rebuild the local site's data and asset copies."""
from pathlib import Path
from datetime import datetime
import json, re, shutil, hashlib
import openpyxl
from PIL import Image, ImageOps

PROJECT = Path(__file__).resolve().parent
ROOT = PROJECT.parent
OUT = PROJECT / 'dist'
ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)

def make_thumbnail(source, destination, max_edge):
    """Create a display-only WebP; keep the JPEG source byte-for-byte."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as original:
        image = ImageOps.exif_transpose(original)
        image.thumbnail((max_edge, max_edge), Image.Resampling.LANCZOS)
        if image.mode not in ('RGB', 'RGBA'):
            image = image.convert('RGB')
        image.save(destination, 'WEBP', quality=85, method=4)
source_txt = ROOT / '影像或病理報告/無影像檔_影像或病理報告_2025-2026.txt'
source_md = ROOT / '元氣診所病歷/病歷轉錄.md'
source_xlsx = next((ROOT / '血液檢驗資料').glob('*.xlsx'))
docs = ASSETS / 'documents'
docs.mkdir(exist_ok=True)
for src in [source_txt, source_md, source_xlsx]:
    shutil.copy2(src, docs / src.name)

reports = []
for block in re.split(r'\n(?=\d{2}\. )', source_txt.read_text())[1:]:
    heading, rest = block.split('\n', 1)
    metadata, body = rest.split('影像或病理報告內容：\n', 1)
    reports.append({'id': heading[:2], 'heading': heading, 'metadata': metadata.strip(), 'body': body.strip()})
groups = []
for r in reports:
    existing = next((g for g in groups if g['body'] == r['body']), None)
    if existing:
        existing['ids'].append(r['id'])
    else:
        groups.append({'title':r['heading'][4:], 'body':r['body'], 'ids':[r['id']]})
display_titles={
    '01':'2026/09/04 · 腰椎 MRI',
    '02':'2026/09/04 · 腰椎屈伸 X 光／KUB 聯合報告',
    '05':'2026/03 · 腸黏膜病理與免疫染色',
    '10':'2025/12 · 右下肢動脈超音波',
    '11':'2025/12 · 右下肢靜脈超音波',
    '14':'2025/12 · 完整 NCV、F 波與 H 反射',
    '22':'2025/08/07 · 腰椎 MRI'
}
for group in groups: group['displayTitle']=display_titles[group['ids'][0]]

w = openpyxl.load_workbook(source_xlsx, data_only=True, read_only=True)
ws = w['完整明細']
headers = [c.value for c in ws[7]]
labs = []
for idx, row in enumerate(ws.iter_rows(min_row=8, values_only=True), 8):
    vals = [v.strftime('%Y-%m-%d') if isinstance(v,datetime) else v for v in row]
    entry = dict(zip(headers, vals))
    entry['cellRange'] = f'A{idx}:L{idx}'
    result, ref = vals[2], vals[4]
    bounds = re.findall(r'\[([^\]]*)\]',ref or '')
    flag = ''
    if isinstance(result,(int,float)) and len(bounds)==2:
        try:
            lo,hi = float(bounds[0]),float(bounds[1])
            flag = '低於原參考值' if result<lo else '高於原參考值' if result>hi else ''
        except ValueError: pass
    entry['rangeFlag'] = flag
    labs.append(entry)

series = []
for year in ['2025','2026']:
    base = ROOT/'MRI 影像'/year
    for folder in sorted({p.parent for p in base.rglob('*.jpg')}):
        label = str(folder.relative_to(base))
        images=[]
        for src in sorted(folder.glob('*.jpg')):
            rel=Path('assets')/'imaging'/year/folder.relative_to(base)/src.name
            dest=OUT/rel
            dest.parent.mkdir(parents=True,exist_ok=True)
            shutil.copy2(src,dest)
            thumb=rel.with_suffix('.webp')
            make_thumbnail(src, OUT/thumb, 1200)
            images.append({'path':rel.as_posix(),'thumb':thumb.as_posix(),'filename':src.name,'original':str(src.relative_to(ROOT))})
        series.append({'id':f'{year}-{len(series)}','year':year,'name':label,'images':images})

photos=[]
for src in sorted((ROOT/'元氣診所病歷/原始照片').glob('*.HEIC')):
    shutil.copy2(src, ASSETS/'records'/src.name)
    jpeg=ASSETS/'records'/f'{src.stem}.jpg'
    if not jpeg.exists():
        raise FileNotFoundError(f'Missing JPEG display copy for {src.name}: {jpeg}')
    thumb=ASSETS/'records'/f'{src.stem}.webp'
    make_thumbnail(jpeg, thumb, 900)
    photos.append({'name':src.stem,'path':f'assets/records/{src.stem}.jpg','thumb':f'assets/records/{src.stem}.webp','original':f'assets/records/{src.name}'})

research=[]
for key in ['orthopedics','rehabilitation','neurosurgery','neurology']:
    path=PROJECT/'research'/f'{key}.json'
    if path.exists():
        data=json.loads(path.read_text());data['id']=key;research.append(data)

clinical=json.loads((PROJECT/'clinical.json').read_text())
data={**clinical,'reports':reports,'reportGroups':groups,'recordText':source_md.read_text(),
      'labs':labs,'labHeaders':headers,'series':series,'photos':photos,'team':research,
      'documents': [{'name':src.name,'path':f'assets/documents/{src.name}'} for src in [source_txt,source_md,source_xlsx]],
      'metadata': {year:(ROOT/'MRI 影像'/year/'MetaData').read_text() for year in ['2025','2026']},
      'counts':{'images':sum(len(s['images']) for s in series),'series':len(series),'reports':len(reports),'reportGroups':len(groups),'labs':len(labs)}}
(OUT/'data.js').write_text('window.CASE_DATA = '+json.dumps(data,ensure_ascii=False)+';\n')
manifest=[]
for p in sorted(ROOT.rglob('*')):
    if p.is_file() and PROJECT not in p.parents and p.name != '.DS_Store' and '.git' not in p.parts:
        manifest.append({'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(PROJECT/'verification/source-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(data['counts'],ensure_ascii=False))
